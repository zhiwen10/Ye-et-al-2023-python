"""Generate the simple 2AFC literature overview (PNG diagram + Word document).

Run with the spirals-py environment:
    python tools/make_reading_map.py

Outputs:
    docs/2afc_reading_map_diagram.png
    docs/2afc_reading_map.docx
"""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

OUT_DIR = Path(__file__).resolve().parent.parent / "docs"
OUT_DIR.mkdir(exist_ok=True)
PNG_PATH = OUT_DIR / "2afc_reading_map_diagram.png"
DOCX_PATH = OUT_DIR / "2afc_reading_map.docx"

# ---------------------------------------------------------------------------
# Content (shared by diagram and document)
# ---------------------------------------------------------------------------

DIRECTIONS = [
    {
        "name": "1. Cortex lab \u2192 IBL",
        "color": "#dbe9f6",
        "task": "Task: mouse sees a faint grating on the left or right, turns a wheel to report.",
        "ask": "Asks: where in the brain are stimulus, choice and action coded?",
        "found": "Found: choice and action signals are brain-wide \u2014 there is no single "
                 "\u201cdecision area\u201d; even the animal\u2019s prior bias is spread across the brain.",
        "papers": [
            ("Burgess 2017", "the wheel task itself"),
            ("Steinmetz 2019", "electrophysiology across the whole brain: choice/action are everywhere"),
            ("IBL 2025", "same task standardized across labs; brain-wide activity map (139 mice)"),
            ("Findling / IBL 2025", "prior expectations (bias) are brain-wide too"),
        ],
    },
    {
        "name": "2. Svoboda lab",
        "color": "#fdebd0",
        "task": "Task: whisker touches a pole; after a short delay, mouse licks left or right.",
        "ask": "Asks: how does the brain hold a movement plan in memory and then execute it?",
        "found": "Found: frontal cortex (ALM) holds the plan as persistent, attractor-like activity.",
        "papers": [
            ("Guo 2014", "the task; ALM is causally necessary"),
            ("Li 2015", "ALM persistent activity encodes the planned choice"),
            ("Li 2016", "the choice signal is robust at the population level"),
            ("Inagaki 2019", "discrete attractor dynamics hold the choice"),
        ],
    },
    {
        "name": "3. Brody lab",
        "color": "#d5f5e3",
        "task": "Task: rat/mouse accumulates streams of clicks (or visual pulses), then orients.",
        "ask": "Asks: how is noisy evidence integrated over time into a decision?",
        "found": "Found: frontal cortex (FOF, the analog of ALM) is required for accumulation; "
                 "parietal cortex is not.",
        "papers": [
            ("Erlich 2011", "FOF: the rat analog of ALM \u2014 links to the Svoboda line"),
            ("Brunton 2013", "the evidence-accumulation task; animals integrate near-optimally"),
            ("Hanks 2015", "FOF is required for accumulation, parietal cortex is not"),
            ("Pinto 2019", "a virtual-reality version for mice, enabling imaging"),
        ],
    },
]

FOOTER = ("Three windows on the same computation \u2014 fast detection, planned action, "
          "evidence accumulation. Ye et al. 2023 (this repository) images the wheel task "
          "cortex-wide and asks how spiral waves fit in.")

# ---------------------------------------------------------------------------
# Diagram
# ---------------------------------------------------------------------------


def make_diagram():
    fig, ax = plt.subplots(figsize=(11, 6.0))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")

    ax.text(50, 96.5, "Three ways to study decision-making in the mouse brain",
            ha="center", fontsize=14, fontweight="bold")
    ax.text(50, 90.5, "one question \u2014 how does sensory evidence become a choice? "
            "\u2014 three task designs",
            ha="center", fontsize=9.5, style="italic")

    y0, h, gap = 59, 25, 4.5
    for i, d in enumerate(DIRECTIONS):
        y = y0 - i * (h + gap)
        p = FancyBboxPatch((2, y), 96, h, boxstyle="round,pad=0.3,rounding_size=1.0",
                           linewidth=1.0, edgecolor="#555555", facecolor=d["color"],
                           zorder=1)
        ax.add_patch(p)
        ax.text(5, y + h - 4.2, d["name"], fontsize=11.5, fontweight="bold", zorder=3)
        ax.text(5, y + h - 9.0, d["task"], fontsize=8.6, zorder=3)
        ax.text(5, y + h - 12.6, d["ask"], fontsize=8.6, zorder=3)
        ax.text(5, y + h - 16.2, d["found"], fontsize=8.6, zorder=3,
                wrap=True)
        ax.text(60, y + h - 4.2, "Example papers", fontsize=9, fontweight="bold",
                zorder=3)
        for j, (cite, line) in enumerate(d["papers"]):
            ax.text(60, y + h - 8.2 - j * 4.6, f"{cite}", fontsize=8.4,
                    fontweight="bold", zorder=3)
            ax.text(60, y + h - 10.6 - j * 4.6, line, fontsize=7.8, zorder=3)

    ax.text(50, 1.5, FOOTER, ha="center", fontsize=8.2, style="italic")

    fig.savefig(PNG_PATH, dpi=200, bbox_inches="tight")
    plt.close(fig)
    print(f"wrote {PNG_PATH}")


# ---------------------------------------------------------------------------
# Word document
# ---------------------------------------------------------------------------


def make_docx():
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    doc.add_heading("Three ways to study decision-making in the mouse brain", level=0)
    doc.add_paragraph(
        "One question — how does sensory evidence become a choice? — studied with three "
        "different task designs. For each: the task, what it asks, what was found, and a "
        "few example papers."
    )

    doc.add_picture(str(PNG_PATH), width=Inches(6.5))
    doc.paragraphs[-1].alignment = 1  # center

    for d in DIRECTIONS:
        doc.add_heading(d["name"], level=1)
        for key in ("task", "ask", "found"):
            p = doc.add_paragraph(d[key])
            p.runs[0].font.size = Pt(10.5)
        for cite, line in d["papers"]:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(f"{cite}: ")
            run.bold = True
            run.font.size = Pt(10)
            run2 = p.add_run(line)
            run2.font.size = Pt(10)

    doc.add_heading("How they fit together", level=1)
    doc.add_paragraph(FOOTER).runs[0].font.size = Pt(10)

    doc.save(DOCX_PATH)
    print(f"wrote {DOCX_PATH}")


if __name__ == "__main__":
    make_diagram()
    make_docx()
